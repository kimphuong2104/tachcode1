# -*- coding: utf-8 -*-

import os
import shutil
import tempfile

import lxml.etree as et  # nosec
import mock
import openpyxl as xl
import pytest

from cdb import testcase
from cs.tools.powerreports.xmlreportgenerator import tools, xsd_util

WORKING_DIR = os.path.dirname(__file__)
TEST_DIR = os.path.join(
    WORKING_DIR, "..", "..", "..", "..", "..", "tests", "test_files"
)


class XSDUtilTestCase(testcase.RollbackTestCase):
    def setUp(self):
        super(XSDUtilTestCase, self).setUp()
        self.test_dir = tempfile.mkdtemp()
        shutil.copy2(os.path.join(TEST_DIR, "empty.xlsx"), self.test_dir)
        shutil.copy2(
            os.path.join(TEST_DIR, "Anforderungsübersicht_template.xlsx"), self.test_dir
        )
        shutil.copy2(
            os.path.join(TEST_DIR, "VarStücklistenvergleich_2_schema.xlsm"),
            self.test_dir,
        )

    @mock.patch.object(xsd_util.XSDUtil, "add")
    def test_empty_excel_add_called(self, mock):
        excel_file = os.path.join(self.test_dir, "empty.xlsx")
        schema = os.path.join(TEST_DIR, "xsd_schema.xsd")
        t = xsd_util.XSDUtil(excel_file, schema, "CDB_RequirementOverview")

        t.import_schema()
        mock.assert_called()

    @mock.patch.object(xsd_util.XSDUtil, "reload_xml_schema")
    def test_req_reloaded_called(self, mock):
        excel_file = os.path.join(self.test_dir, "Anforderungsübersicht_template.xlsx")
        schema = os.path.join(TEST_DIR, "xsd_schema.xsd")
        t = xsd_util.XSDUtil(excel_file, schema, "CDB_RequirementOverview")

        t.import_schema()
        mock.assert_called()

    @pytest.mark.xfail(raises=AssertionError)
    @mock.patch.object(xsd_util.XSDUtil, "add")
    def test_req_fail_add_called(self, mock):
        excel_file = os.path.join(self.test_dir, "Anforderungsübersicht_template.xlsx")
        schema = os.path.join(TEST_DIR, "xsd_schema.xsd")
        t = xsd_util.XSDUtil(excel_file, schema, "CDB_RequirementOverview")

        t.import_schema()
        mock.assert_called()

    @pytest.mark.xfail(raises=AssertionError)
    @mock.patch.object(xsd_util.XSDUtil, "reload_xml_schema")
    def test_empty_excel_fail_reloaded_called(self, mock):
        excel_file = os.path.join(self.test_dir, "empty.xlsx")
        schema = os.path.join(TEST_DIR, "xsd_schema.xsd")
        t = xsd_util.XSDUtil(excel_file, schema, "CDB_RequirementOverview")

        t.import_schema()
        mock.assert_called()

    def test_add_xsd_of_req_overview_to_empty(self):
        excel_file = os.path.join(self.test_dir, "empty.xlsx")
        schema = os.path.join(TEST_DIR, "xsd_schema.xsd")
        client_excel = os.path.join(TEST_DIR, "Anforderungsübersicht_caddok.xlsx")

        xsd_util.XSDUtil(excel_file, schema, "CDB_RequirementOverview").add()

        excel_child_map = self.get_tree_child_map(excel_file)
        client_excel_child_map = self.get_tree_child_map(client_excel)

        # check if xmlMaps have same RootTypes (by name)
        # .getchildren()[0] for <xsd:all>
        for idx in range(0, 4):
            excel_element = (
                excel_child_map.getchildren()[0].getchildren()[idx].attrib["name"]
            )
            client_element = (
                client_excel_child_map.getchildren()[0]
                .getchildren()[idx]
                .attrib["name"]
            )
            self.assertEqual(excel_element, client_element)

    def test_reload_xml_only_cdb(self):
        excel_file = os.path.join(
            self.test_dir, "VarStücklistenvergleich_2_schema.xlsm"
        )
        schema = os.path.join(TEST_DIR, "xsd_schema_new.xsd")
        xsd_util.XSDUtil(
            excel_file, schema, "CDB_RequirementOverview"
        ).reload_xml_schema()
        temp_dir = tools.temporary_unzip_file(excel_file)

        map_file = os.path.join(temp_dir, "xl", "xmlMaps.xml")
        map_root = et.parse(map_file).getroot()  # pylint: disable=I1101 #nosec
        schemas = []
        expected_schemas = ["CDB_RequirementOverview", "Root_Zuordnung"]
        number_schemas = 0
        expected_number_schemas = 2
        for child in map_root:
            if "Schema" in child.tag:
                number_schemas += 1
            if "Map" in child.tag:
                schemas.append(child.attrib["Name"])

        shutil.rmtree(temp_dir)

        self.assertEqual(schemas.sort(), expected_schemas.sort())
        self.assertEqual(number_schemas, expected_number_schemas)

    def test_reload_xml_new(self):
        excel_file = os.path.join(self.test_dir, "Anforderungsübersicht_template.xlsx")
        schema = os.path.join(TEST_DIR, "xsd_schema_new.xsd")

        workbook = xl.load_workbook(excel_file)
        sheet1 = workbook.worksheets[0]
        for table in sheet1.tables.values():
            for table_column in table.tableColumns:
                if table_column.uniqueName == "chapter":
                    chapter_column_old_xml = table_column.xmlColumnPr
                if table_column.uniqueName == "act_value":
                    act_value_column_old_xml = table_column.xmlColumnPr

        xsd_util.XSDUtil(
            excel_file, schema, "CDB_RequirementOverview"
        ).reload_xml_schema()
        temp_dir = tools.temporary_unzip_file(excel_file)
        map_file = os.path.join(temp_dir, "xl", "xmlMaps.xml")
        map_root = et.parse(map_file).getroot()  # pylint: disable=I1101 #nosec
        number_schemas = 0
        expected_number_schemas = 1
        for child in map_root:
            if "Schema" in child.tag:
                number_schemas += 1

        self.assertEqual(number_schemas, expected_number_schemas)

        workbook = xl.load_workbook(excel_file)
        sheet1 = workbook.worksheets[0]
        for table in sheet1.tables.values():
            for table_column in table.tableColumns:
                if table_column.uniqueName == "chapter":
                    chapter_column_new_xml = table_column.xmlColumnPr
                if table_column.uniqueName == "act_value":
                    act_value_column_new_xml = table_column.xmlColumnPr

        self.assertEqual(chapter_column_new_xml, chapter_column_old_xml)
        self.assertEqual(act_value_column_new_xml, act_value_column_old_xml)

    @staticmethod
    def get_tree_child_map(excel):
        temp_dir = tools.temporary_unzip_file(excel)

        map_file = os.path.join(temp_dir, "xl", "xmlMaps.xml")
        map_root = et.parse(map_file).getroot()  # pylint: disable=I1101 #nosec

        child_map = ""

        for child in map_root:
            for child1 in child:
                if child1.tag == "{http://www.w3.org/2001/XMLSchema}schema":
                    child_map = child1.getchildren()[1]

        shutil.rmtree(temp_dir)

        return child_map

    def tearDown(self):
        super(XSDUtilTestCase, self).tearDown()
        shutil.rmtree(self.test_dir)
